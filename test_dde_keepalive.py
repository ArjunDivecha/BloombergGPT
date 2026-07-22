"""
=============================================================================
SCRIPT NAME: test_dde_keepalive.py
=============================================================================

INPUT FILES:
- None (connects to Bloomberg Terminal via Parallels Windows VM)

OUTPUT FILES:
- outputs/dde_keepalive_log_YYYY_MM_DD_HHMMSS.xlsx: Full log of every probe
- outputs/dde_keepalive_outages_YYYY_MM_DD_HHMMSS.csv: Outage windows
- outputs/dde_keepalive_summary.txt: Run summary appended at end

VERSION: 2.0
LAST UPDATED: 2026-02-24
AUTHOR: Arjun Divecha

DESCRIPTION:
Long-running Bloomberg connection investigator. This script is designed for
multi-day experiments and keeps running through failures/sleep/wake. It tracks:
1. Continuous UP and DOWN windows (with start/end/duration)
2. Why failures happened (session, network/API port, process checks)
3. Sleep-gap events (when machine likely slept between loops)
4. Recovery timing when Bloomberg comes back

DEPENDENCIES:
- blpapi (installed via conda-forge in OpusBloomberg .venv)
- openpyxl (for xlsx output, optional)

USAGE:
    "/Users/arjundivecha/Dropbox/AAA Backup/A Working/OpusBloomberg/.venv/bin/python" \
        test_dde_keepalive.py --interval 300 --max-hours 120

    Optional arguments:
        --interval 300         Seconds between normal pings (default: 300 = 5 min)
        --retry-interval 60    Seconds between retries while DOWN (default: 60)
        --max-hours 120        Max runtime in hours; set 0 for no max
        --no-process-probe     Disable Windows process polling via prlctl/tasklist

NOTES:
- Bloomberg Terminal should be open/logged in on Windows side for UP state.
- Results are written after every probe (fault-tolerant).
- Press Ctrl+C to stop gracefully; summary and outage report are still written.
=============================================================================
"""

import argparse
import csv
import os
import signal
import socket
import subprocess
import sys
import time
from datetime import datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

# Add OpusBloomberg to path
sys.path.insert(0, "/Users/arjundivecha/Dropbox/AAA Backup/A Working/OpusBloomberg")

from bbg import BBG, detect_vm_ip

# =============================================================================
# CONFIGURATION
# =============================================================================

TICKER_ROTATION = [
    ("AAPL US Equity", ["PX_LAST", "NAME"]),
    ("MSFT US Equity", ["PX_LAST", "PE_RATIO"]),
    ("SPX Index", ["PX_LAST", "CHG_PCT_1D"]),
    ("USDJPY Curncy", ["PX_LAST", "NAME"]),
    ("GT10 Govt", ["PX_LAST", "NAME"]),
    ("GOOGL US Equity", ["PX_LAST", "CUR_MKT_CAP"]),
    ("CL1 Comdty", ["PX_LAST", "NAME"]),
    ("GC1 Comdty", ["PX_LAST", "NAME"]),
    ("NKY Index", ["PX_LAST", "NAME"]),
    ("EURUSD Curncy", ["PX_LAST", "NAME"]),
]

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outputs")

VM_NAME = "Windows 11"
BLOOMBERG_PORT = 8194
TERMINAL_PROCESS_NAMES = ["terminal.exe", "wintrv.exe"]
BBCOMM_PROCESS_NAME = "bbcomm.exe"


# =============================================================================
# HELPERS
# =============================================================================

def format_duration(seconds: float) -> str:
    """Render seconds as HH:MM:SS."""
    return str(timedelta(seconds=int(max(seconds, 0))))


def bool_to_text(value: Optional[bool]) -> str:
    """Convert tri-state bool to display text."""
    if value is True:
        return "YES"
    if value is False:
        return "NO"
    return "UNKNOWN"


def check_api_port(host: str, port: int = BLOOMBERG_PORT, timeout: float = 3.0) -> Tuple[bool, str]:
    """Check whether Bloomberg API TCP port is reachable."""
    if not host:
        return False, "No host provided"

    try:
        sock = socket.create_connection((host, port), timeout=timeout)
        sock.close()
        return True, ""
    except Exception as exc:
        return False, str(exc)


def probe_bloomberg_processes(vm_name: str = VM_NAME) -> Dict[str, Optional[bool]]:
    """
    Poll Windows VM process list for Bloomberg client processes.
    This does not guarantee API usability, but it gives useful context.
    """
    try:
        result = subprocess.run(
            ["prlctl", "exec", vm_name, "cmd.exe", "/c", "tasklist"],
            capture_output=True,
            text=True,
            timeout=20,
        )
        raw = (result.stdout or "").lower()
        terminal_running = any(name in raw for name in TERMINAL_PROCESS_NAMES)
        bbcomm_running = BBCOMM_PROCESS_NAME in raw
        return {
            "terminal_running": terminal_running,
            "bbcomm_running": bbcomm_running,
            "error": None,
        }
    except Exception as exc:
        return {
            "terminal_running": None,
            "bbcomm_running": None,
            "error": str(exc),
        }


def classify_failure(
    error: str,
    vm_ip_error: Optional[str],
    api_port_ok: Optional[bool],
    terminal_running: Optional[bool],
    bbcomm_running: Optional[bool],
) -> str:
    """Classify failures for easier post-run analysis."""
    if vm_ip_error:
        return "VM_IP_UNAVAILABLE"

    msg = (error or "").lower()
    if "session not started" in msg:
        return "SESSION_NOT_STARTED"
    if "failed to start bloomberg session" in msg:
        return "SESSION_START_FAILED"
    if "failed to open //blp/refdata" in msg:
        return "REFDATA_SERVICE_UNAVAILABLE"
    if "timed out" in msg:
        return "TIMEOUT"

    if api_port_ok is False:
        return "API_PORT_UNREACHABLE"
    if bbcomm_running is False and terminal_running is False:
        return "TERMINAL_AND_BBCOMM_NOT_RUNNING"
    if bbcomm_running is False:
        return "BBCOMM_NOT_RUNNING"
    if terminal_running is False:
        return "TERMINAL_NOT_RUNNING"
    return "UNKNOWN_FAILURE"


# =============================================================================
# RESULT LOGGING
# =============================================================================

def init_results_file() -> str:
    """Create results workbook/csv and return file path."""
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    timestamp = datetime.now().strftime("%Y_%m_%d_%H%M%S")
    filepath = os.path.join(OUTPUT_DIR, f"dde_keepalive_log_{timestamp}.xlsx")

    headers = [
        "Probe #",
        "Timestamp",
        "Elapsed (HH:MM:SS)",
        "Elapsed (seconds)",
        "State",
        "Ticker",
        "Fields Requested",
        "Data Received",
        "Response Time (ms)",
        "Error Message",
        "Failure Category",
        "VM IP",
        "API Port Reachable",
        "Terminal Process Running",
        "BBCOMM Process Running",
        "Sleep Gap Detected (s)",
        "Recovery Action",
        "Cumulative Successes",
        "Cumulative Failures",
        "Success Rate (%)",
        "Total Downtime (HH:MM:SS)",
        "Outage Count",
    ]

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Keepalive Investigator"
        ws.append(headers)
        bold = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold

        widths = {
            "A": 8,
            "B": 22,
            "C": 15,
            "D": 15,
            "E": 8,
            "F": 20,
            "G": 28,
            "H": 44,
            "I": 18,
            "J": 44,
            "K": 24,
            "L": 16,
            "M": 18,
            "N": 24,
            "O": 24,
            "P": 20,
            "Q": 28,
            "R": 20,
            "S": 20,
            "T": 16,
            "U": 22,
            "V": 14,
        }
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        wb.save(filepath)
    except ImportError:
        filepath = filepath.replace(".xlsx", ".csv")
        with open(filepath, "w", newline="") as file_obj:
            csv.writer(file_obj).writerow(headers)

    return filepath


def append_result(filepath: str, row_data: List[Any]) -> None:
    """Append one row to xlsx/csv results file using an atomic xlsx write."""
    if filepath.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(filepath)
        ws = wb.active
        ws.append(row_data)
        tmp = filepath + ".tmp"
        wb.save(tmp)
        os.replace(tmp, filepath)
    else:
        with open(filepath, "a", newline="") as file_obj:
            csv.writer(file_obj).writerow(row_data)


def write_outage_report(filepath: str, outages: List[Dict[str, Any]]) -> None:
    """Write outage windows to a standalone csv report."""
    with open(filepath, "w", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(
            ["Outage #", "Start", "End", "Duration (HH:MM:SS)", "Duration (seconds)", "Cause"]
        )
        for idx, outage in enumerate(outages, start=1):
            duration_sec = float(outage["duration_sec"])
            writer.writerow(
                [
                    idx,
                    outage["start"],
                    outage["end"],
                    format_duration(duration_sec),
                    round(duration_sec, 1),
                    outage["cause"],
                ]
            )


# =============================================================================
# PROBE EXECUTION
# =============================================================================

def ping_bloomberg(bbg: BBG, ticker: str, fields: List[str]) -> Dict[str, Any]:
    """Pull data from Bloomberg and return success/data/error payload."""
    start = time.perf_counter()
    try:
        data = bbg.ref(ticker, fields)
        elapsed_ms = (time.perf_counter() - start) * 1000
        has_data = any(value is not None for value in data.values())
        if has_data:
            return {
                "success": True,
                "data": data,
                "response_ms": round(elapsed_ms, 1),
                "error": None,
            }
        return {
            "success": False,
            "data": data,
            "response_ms": round(elapsed_ms, 1),
            "error": "All fields returned None",
        }
    except Exception as exc:
        elapsed_ms = (time.perf_counter() - start) * 1000
        return {
            "success": False,
            "data": None,
            "response_ms": round(elapsed_ms, 1),
            "error": str(exc),
        }


def run_keepalive_investigator(
    interval_sec: int = 300,
    retry_interval_sec: int = 60,
    max_hours: float = 120,
    probe_processes: bool = True,
    sleep_gap_multiplier: float = 1.5,
) -> None:
    """
    Long-running investigator loop.

    Args:
        interval_sec: Probe cadence while connection is UP.
        retry_interval_sec: Probe cadence while connection is DOWN.
        max_hours: Max runtime in hours. Set 0 to run indefinitely.
        probe_processes: Poll Windows tasklist for terminal/bbcomm process context.
        sleep_gap_multiplier: Factor used to detect likely machine sleep gaps.
    """
    max_seconds = max_hours * 3600 if max_hours > 0 else None

    print("=" * 76)
    print("  BLOOMBERG DDE KEEPALIVE INVESTIGATOR")
    print(f"  Started: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"  UP interval: {interval_sec}s ({interval_sec / 60:.1f} min)")
    print(f"  DOWN retry interval: {retry_interval_sec}s")
    print(f"  Max duration: {'Unlimited' if max_seconds is None else f'{max_hours} hours'}")
    print(f"  Process probe: {'ON' if probe_processes else 'OFF'}")
    print("=" * 76)
    print()

    results_file = init_results_file()
    timestamp_tag = os.path.basename(results_file).split("dde_keepalive_log_")[-1].split(".")[0]
    outages_file = os.path.join(OUTPUT_DIR, f"dde_keepalive_outages_{timestamp_tag}.csv")
    print(f"Results file: {results_file}")
    print(f"Outages file: {outages_file}")
    print()

    start_time = datetime.now()
    ping_num = 0
    successes = 0
    failures = 0
    first_failure_time: Optional[datetime] = None
    last_success_time: Optional[datetime] = None

    bbg: Optional[BBG] = None
    state = "UNKNOWN"
    current_uptime_start: Optional[datetime] = None
    current_downtime_start: Optional[datetime] = None
    current_downtime_cause: Optional[str] = None
    uptime_windows_sec: List[float] = []
    outages: List[Dict[str, Any]] = []
    total_downtime_sec = 0.0
    sleep_gap_events = 0

    last_loop_started: Optional[datetime] = None
    last_expected_wait = interval_sec
    shutdown = False

    def handle_signal(signum: int, frame: Any) -> None:
        del signum, frame
        nonlocal shutdown
        shutdown = True
        print("\n\nCtrl+C received. Finishing current probe and writing summary...")

    signal.signal(signal.SIGINT, handle_signal)

    print("=" * 76)
    print(
        f"  {'#':>4}  {'Time':>10}  {'Elapsed':>12}  "
        f"{'State':>6}  {'ms':>7}  {'Gap':>6}  Probe"
    )
    print("-" * 76)

    try:
        while not shutdown:
            now = datetime.now()
            elapsed_sec = (now - start_time).total_seconds()

            if max_seconds is not None and elapsed_sec >= max_seconds:
                print(f"\nMax duration of {max_hours} hours reached. Stopping.")
                break

            sleep_gap_sec = 0.0
            if last_loop_started is not None:
                observed_gap = (now - last_loop_started).total_seconds()
                threshold = max(20.0, last_expected_wait * sleep_gap_multiplier)
                if observed_gap > threshold:
                    sleep_gap_sec = round(observed_gap - last_expected_wait, 1)
                    sleep_gap_events += 1
                    print(
                        f"\n  >>> Sleep/hibernation gap detected: +{sleep_gap_sec}s "
                        f"(observed {observed_gap:.1f}s between probes)"
                    )
            last_loop_started = now

            ping_num += 1
            ticker, fields = TICKER_ROTATION[(ping_num - 1) % len(TICKER_ROTATION)]
            recovery_actions: List[str] = []

            vm_ip: Optional[str] = None
            vm_ip_error: Optional[str] = None
            try:
                vm_ip = detect_vm_ip()
            except Exception as exc:
                vm_ip_error = str(exc)

            api_port_ok: Optional[bool] = None
            api_port_error = ""
            if vm_ip:
                api_port_ok, api_port_error = check_api_port(vm_ip, BLOOMBERG_PORT)

            process_probe = {"terminal_running": None, "bbcomm_running": None, "error": None}
            if probe_processes:
                process_probe = probe_bloomberg_processes()

            terminal_running = process_probe["terminal_running"]
            bbcomm_running = process_probe["bbcomm_running"]

            if bbg is not None and vm_ip and bbg.host != vm_ip:
                try:
                    bbg.disconnect()
                except Exception:
                    pass
                bbg = None
                recovery_actions.append(f"HOST_CHANGED->{vm_ip}")

            connect_error = None
            if bbg is None and vm_ip and api_port_ok is not False:
                try:
                    bbg = BBG(host=vm_ip)
                    bbg.connect()
                    recovery_actions.append("CONNECTED")
                except Exception as exc:
                    bbg = None
                    connect_error = str(exc)
                    recovery_actions.append("CONNECT_FAILED")

            if bbg is not None:
                result = ping_bloomberg(bbg, ticker, fields)
            else:
                reason_bits = []
                if vm_ip_error:
                    reason_bits.append(f"VM IP detect failed: {vm_ip_error}")
                if connect_error:
                    reason_bits.append(f"Connect failed: {connect_error}")
                if vm_ip and api_port_ok is False:
                    reason_bits.append(
                        f"API port {BLOOMBERG_PORT} unreachable: {api_port_error}"
                    )
                if not reason_bits:
                    reason_bits.append("No active Bloomberg session")
                result = {
                    "success": False,
                    "data": None,
                    "response_ms": 0.0,
                    "error": " | ".join(reason_bits),
                }

            if result["success"]:
                successes += 1
                last_success_time = now
            else:
                failures += 1
                if first_failure_time is None:
                    first_failure_time = now
                if bbg is not None:
                    try:
                        bbg.disconnect()
                    except Exception:
                        pass
                    bbg = None
                    recovery_actions.append("SESSION_RESET")

            failure_category = ""
            if not result["success"]:
                failure_category = classify_failure(
                    error=result["error"] or "",
                    vm_ip_error=vm_ip_error,
                    api_port_ok=api_port_ok,
                    terminal_running=terminal_running,
                    bbcomm_running=bbcomm_running,
                )

            new_state = "UP" if result["success"] else "DOWN"
            if new_state != state:
                if new_state == "UP":
                    if current_downtime_start is not None:
                        outage_sec = (now - current_downtime_start).total_seconds()
                        total_downtime_sec += outage_sec
                        outages.append(
                            {
                                "start": current_downtime_start.strftime("%Y-%m-%d %H:%M:%S"),
                                "end": now.strftime("%Y-%m-%d %H:%M:%S"),
                                "duration_sec": outage_sec,
                                "cause": current_downtime_cause or "UNKNOWN_FAILURE",
                            }
                        )
                        print(
                            f"        RECOVERED after {format_duration(outage_sec)} "
                            f"(cause: {current_downtime_cause or 'UNKNOWN_FAILURE'})"
                        )
                    current_downtime_start = None
                    current_downtime_cause = None
                    current_uptime_start = now
                else:
                    if current_uptime_start is not None:
                        uptime_sec = (now - current_uptime_start).total_seconds()
                        uptime_windows_sec.append(uptime_sec)
                    current_uptime_start = None
                    current_downtime_start = now
                    current_downtime_cause = failure_category or "UNKNOWN_FAILURE"
                    print(
                        f"        OUTAGE started (reason: {current_downtime_cause})"
                    )
                state = new_state

            total = successes + failures
            success_rate = (successes / total * 100) if total > 0 else 0.0
            elapsed_str = format_duration(elapsed_sec)

            if result["data"]:
                data_str = str(result["data"])
                if len(data_str) > 62:
                    data_str = data_str[:59] + "..."
            else:
                data_str = "N/A"

            active_downtime_sec = (
                (now - current_downtime_start).total_seconds()
                if current_downtime_start is not None
                else 0.0
            )
            total_downtime_for_row = total_downtime_sec + active_downtime_sec
            outage_count = len(outages) + (1 if current_downtime_start is not None else 0)

            print(
                f"  {ping_num:>4}  {now.strftime('%H:%M:%S'):>10}  {elapsed_str:>12}  "
                f"{new_state:>6}  {result['response_ms']:>6.0f}  "
                f"{sleep_gap_sec:>5.0f}s  {ticker}: {data_str}"
            )
            if result["error"]:
                print(f"        ERROR: {result['error']}")

            row = [
                ping_num,
                now.strftime("%Y-%m-%d %H:%M:%S"),
                elapsed_str,
                int(elapsed_sec),
                new_state,
                ticker,
                ", ".join(fields),
                str(result["data"]) if result["data"] else "",
                result["response_ms"],
                result["error"] or "",
                failure_category,
                vm_ip or "",
                bool_to_text(api_port_ok),
                bool_to_text(terminal_running),
                bool_to_text(bbcomm_running),
                sleep_gap_sec,
                "; ".join(recovery_actions),
                successes,
                failures,
                round(success_rate, 1),
                format_duration(total_downtime_for_row),
                outage_count,
            ]

            try:
                append_result(results_file, row)
            except Exception as exc:
                print(f"        WARN: Could not write to results file: {exc}")

            last_expected_wait = interval_sec if result["success"] else retry_interval_sec
            wait_until = time.time() + last_expected_wait
            while time.time() < wait_until and not shutdown:
                time.sleep(1)

    finally:
        end_time = datetime.now()
        total_elapsed_sec = (end_time - start_time).total_seconds()

        if state == "UP" and current_uptime_start is not None:
            uptime_windows_sec.append((end_time - current_uptime_start).total_seconds())
        if state == "DOWN" and current_downtime_start is not None:
            outage_sec = (end_time - current_downtime_start).total_seconds()
            total_downtime_sec += outage_sec
            outages.append(
                {
                    "start": current_downtime_start.strftime("%Y-%m-%d %H:%M:%S"),
                    "end": "ONGOING_AT_STOP",
                    "duration_sec": outage_sec,
                    "cause": current_downtime_cause or "UNKNOWN_FAILURE",
                }
            )

        total = successes + failures
        longest_uptime = max(uptime_windows_sec) if uptime_windows_sec else 0.0
        longest_outage = max((float(item["duration_sec"]) for item in outages), default=0.0)
        availability = (
            (max(total_elapsed_sec - total_downtime_sec, 0.0) / total_elapsed_sec * 100)
            if total_elapsed_sec > 0
            else 0.0
        )

        print()
        print("=" * 76)
        print("  KEEPALIVE INVESTIGATION SUMMARY")
        print("=" * 76)
        print(f"  Start time:             {start_time.strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"  End time:               {end_time.strftime('%Y-%m-%d %H:%M:%S')}")
        print(f"  Total runtime:          {format_duration(total_elapsed_sec)}")
        print(f"  Total probes:           {total}")
        print(f"  Successful probes:      {successes}")
        print(f"  Failed probes:          {failures}")
        print(f"  Success rate:           {(successes / total * 100):.1f}%" if total else "  Success rate:           N/A")
        print(f"  Availability:           {availability:.2f}%")
        print(f"  Total downtime:         {format_duration(total_downtime_sec)}")
        print(f"  Outage count:           {len(outages)}")
        print(f"  Longest uptime window:  {format_duration(longest_uptime)}")
        print(f"  Longest outage window:  {format_duration(longest_outage)}")
        print(f"  Sleep-gap events:       {sleep_gap_events}")
        print(
            f"  Last successful probe:  "
            f"{last_success_time.strftime('%Y-%m-%d %H:%M:%S') if last_success_time else 'None'}"
        )
        if first_failure_time:
            print(f"  First failure:          {first_failure_time.strftime('%Y-%m-%d %H:%M:%S')}")
            print(
                f"  Time to first failure:  "
                f"{format_duration((first_failure_time - start_time).total_seconds())}"
            )
        else:
            print("  First failure:          None")
        print(f"  Results file:           {results_file}")
        print(f"  Outages file:           {outages_file}")
        print("=" * 76)

        summary_file = os.path.join(OUTPUT_DIR, "dde_keepalive_summary.txt")
        try:
            with open(summary_file, "a", encoding="utf-8") as file_obj:
                file_obj.write(f"\n{'=' * 76}\n")
                file_obj.write(
                    f"KEEPALIVE INVESTIGATOR — {start_time.strftime('%Y-%m-%d %H:%M:%S')}\n"
                )
                file_obj.write(f"{'=' * 76}\n")
                file_obj.write(f"End:               {end_time.strftime('%Y-%m-%d %H:%M:%S')}\n")
                file_obj.write(f"Runtime:           {format_duration(total_elapsed_sec)}\n")
                file_obj.write(f"Probes:            {total} ({successes} UP, {failures} DOWN)\n")
                file_obj.write(
                    f"Success rate:      {(successes / total * 100):.1f}%\n"
                    if total
                    else "Success rate:      N/A\n"
                )
                file_obj.write(f"Availability:      {availability:.2f}%\n")
                file_obj.write(f"Total downtime:    {format_duration(total_downtime_sec)}\n")
                file_obj.write(f"Outage count:      {len(outages)}\n")
                file_obj.write(f"Longest uptime:    {format_duration(longest_uptime)}\n")
                file_obj.write(f"Longest outage:    {format_duration(longest_outage)}\n")
                file_obj.write(f"Sleep-gap events:  {sleep_gap_events}\n")
                file_obj.write(
                    "Last success:      "
                    f"{last_success_time.strftime('%Y-%m-%d %H:%M:%S') if last_success_time else 'None'}\n"
                )
                if first_failure_time:
                    file_obj.write(
                        f"First failure:     {first_failure_time.strftime('%Y-%m-%d %H:%M:%S')}\n"
                    )
                    file_obj.write(
                        "Time to first:     "
                        f"{format_duration((first_failure_time - start_time).total_seconds())}\n"
                    )
                else:
                    file_obj.write("First failure:     None\n")
                file_obj.write(f"Results:           {results_file}\n")
                file_obj.write(f"Outages:           {outages_file}\n")

                if outages:
                    file_obj.write("Outage windows:\n")
                    for idx, outage in enumerate(outages, start=1):
                        file_obj.write(
                            f"  {idx:>2}. {outage['start']} -> {outage['end']} | "
                            f"{format_duration(float(outage['duration_sec']))} | "
                            f"{outage['cause']}\n"
                        )
                else:
                    file_obj.write("Outage windows:    None\n")
        except Exception as exc:
            print(f"WARN: Could not write summary file: {exc}")

        try:
            write_outage_report(outages_file, outages)
        except Exception as exc:
            print(f"WARN: Could not write outage report: {exc}")

        if bbg is not None:
            try:
                bbg.disconnect()
            except Exception:
                pass

        print(f"\nSummary appended to: {summary_file}")
        print("Done.")


# =============================================================================
# ENTRY POINT
# =============================================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Multi-day Bloomberg DDE investigator for uptime/downtime behavior"
    )
    parser.add_argument(
        "--interval",
        type=int,
        default=300,
        help="Seconds between probes when connection is UP (default: 300)",
    )
    parser.add_argument(
        "--retry-interval",
        type=int,
        default=60,
        help="Seconds between probes when connection is DOWN (default: 60)",
    )
    parser.add_argument(
        "--max-hours",
        type=float,
        default=120,
        help="Maximum runtime in hours; set 0 for no limit (default: 120)",
    )
    parser.add_argument(
        "--no-process-probe",
        action="store_true",
        help="Disable VM process polling via prlctl/tasklist",
    )
    parser.add_argument(
        "--sleep-gap-multiplier",
        type=float,
        default=1.5,
        help="Multiplier for detecting sleep/wake gaps relative to expected wait",
    )
    args = parser.parse_args()

    run_keepalive_investigator(
        interval_sec=args.interval,
        retry_interval_sec=args.retry_interval,
        max_hours=args.max_hours,
        probe_processes=not args.no_process_probe,
        sleep_gap_multiplier=max(args.sleep_gap_multiplier, 1.0),
    )
